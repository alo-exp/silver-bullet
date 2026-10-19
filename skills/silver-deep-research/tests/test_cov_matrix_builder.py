"""Coverage tests for matrix_builder (real openpyxl workbooks written to temp dirs)."""

import json
import logging
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

import covhelp
from covhelp import SAMPLE_CONFIG, SCRIPTS, build_sample_matrix, call_main, run_script

import deps_probe
import matrix_builder as mb

HAS_OPENPYXL = deps_probe.openpyxl_importable()


def setUpModule():
    covhelp.block_network()


def tearDownModule():
    covhelp.unblock_network()


class TestSanitize(unittest.TestCase):
    def test_sanitize(self):
        f = mb._sanitize_cell_value
        self.assertEqual(f("=SUM(A1)"), "'=SUM(A1)")
        self.assertEqual(f("+1"), "'+1")
        self.assertEqual(f("-1"), "'-1")
        self.assertEqual(f("@x"), "'@x")
        self.assertEqual(f("  \t=x"), "'  \t=x")  # leading whitespace does not hide a formula
        self.assertEqual(f("safe"), "safe")
        self.assertEqual(f("   "), "   ")
        self.assertEqual(f(""), "")
        self.assertEqual(f(None), None)
        self.assertEqual(f(5), 5)

    def test_formula_text(self):
        self.assertEqual(mb._countif("C"), '=COUNTIF(C5:C1048576,"✔")')
        score = mb._score_formula("C")
        self.assertTrue(score.startswith("=COUNTIFS($B$5:$B$1048576,\"Critical\",C$5:C$1048576,\"✔\")*5+"))
        self.assertTrue(score.endswith('"Low",C$5:C$1048576,"✔")*1'))
        self.assertEqual(score.count("COUNTIFS"), 5)


@unittest.skipUnless(HAS_OPENPYXL, "openpyxl not installed")
class TestOpenpyxlGuard(unittest.TestCase):
    def test_import_error_message(self):
        with mock.patch.object(mb, "openpyxl", None), \
                mock.patch.dict(sys.modules, {"openpyxl": None}):
            with self.assertRaises(ImportError) as ctx:
                mb._require_openpyxl()
        self.assertIn("openpyxl is required for matrix_builder", str(ctx.exception))

    def test_lazy_init_and_second_call(self):
        mb._require_openpyxl()
        self.assertIsNotNone(mb.openpyxl)
        mb._require_openpyxl()  # already loaded: returns immediately
        # re-running the loader with styles already built keeps the existing styles
        thin = mb._THIN
        with mock.patch.object(mb, "openpyxl", None):
            mb._require_openpyxl()
        self.assertIs(mb._THIN, thin)
        self.assertIsNotNone(mb.openpyxl)

    def test_style_helpers(self):
        mb._require_openpyxl()
        wb = mb.openpyxl.Workbook()
        ws = wb.active
        cell = ws.cell(1, 1)
        mb._style_cell(cell)  # nothing to apply
        self.assertFalse(cell.font.bold)
        mb._style_cell(cell, font=mb._TITLE_FONT)
        self.assertTrue(cell.font.bold)
        self.assertEqual(cell.fill.fgColor.rgb, "00000000")
        mb._style_cell(cell, fill=mb._TITLE_FILL)
        self.assertEqual(cell.fill.fgColor.rgb, "001F3864")
        mb._style_cell(cell, align=mb._TITLE_ALIGN)
        self.assertEqual(cell.alignment.horizontal, "center")
        mb._style_cell(cell, border=mb._DEFAULT_BORDER)
        self.assertEqual(cell.border.left.style, "thin")
        c2 = mb._set_user_cell(ws, 2, 1, "=evil()", mb._FEAT_FONT)
        self.assertEqual(c2.value, "'=evil()")


@unittest.skipUnless(HAS_OPENPYXL, "openpyxl not installed")
class TestBuildMatrix(unittest.TestCase):
    def setUp(self):
        self._td = tempfile.TemporaryDirectory()
        self.addCleanup(self._td.cleanup)
        self.dir = Path(self._td.name)

    def load(self, path):
        import openpyxl
        return openpyxl.load_workbook(path)

    def test_layout_and_content(self):
        out = self.dir / "m.xlsx"
        result = build_sample_matrix(out)
        self.assertEqual(result, {"platforms_added": 3, "platform_names": ["Alpha", "Beta", "Gamma"],
                                  "categories": 2, "features": 4, "total_rows": 10, "output": str(out)})
        ws = self.load(out).active
        self.assertEqual(ws.title, "Comparison Matrix")
        self.assertEqual(ws["A1"].value, "Sample Matrix")
        self.assertIn("A1:E1", [str(r) for r in ws.merged_cells.ranges])
        self.assertEqual([ws.cell(2, c).value for c in range(1, 6)],
                         ["Capability / Feature", "Priority", "Alpha", "Beta", "Gamma"])
        self.assertEqual(ws["A3"].value, "Total Capabilities")
        self.assertEqual(ws["C3"].value, mb._countif("C"))
        self.assertEqual(ws["D4"].value, mb._score_formula("D"))
        self.assertEqual(ws["A4"].value, "Score")
        self.assertEqual(ws["B3"].value, "—")
        self.assertEqual(ws["A5"].value, "Sec")
        self.assertIn("A5:E5", [str(r) for r in ws.merged_cells.ranges])
        self.assertEqual((ws["A6"].value, ws["B6"].value), ("SSO", "Critical"))
        self.assertEqual([ws.cell(6, c).value for c in (3, 4, 5)], ["✔", "✔", "—"])
        self.assertEqual([ws.cell(7, c).value for c in (3, 4, 5)], ["—", "✔", "—"])
        self.assertEqual(ws["A8"].value, "Ops")
        self.assertEqual(ws.freeze_panes, "A5")
        self.assertEqual(ws.auto_filter.ref, "A2:E2")
        self.assertFalse(ws.sheet_view.showGridLines)
        self.assertEqual(ws.column_dimensions["A"].width, 50)
        self.assertEqual(ws.column_dimensions["E"].width, 18)
        self.assertEqual(ws["B6"].fill.fgColor.rgb, "00F4C7C3")  # Critical fill

    def test_defaults_and_unknown_priority(self):
        out = self.dir / "d.xlsx"
        cfg = {"categories": [{"name": "C", "features": [{"name": "F1", "priority": "Weird"}, {"name": "F2"}]},
                              {"name": "Empty"}],
               "platforms": [{"name": "P"}]}
        result = mb.build_matrix(cfg, str(out))
        self.assertEqual(result["features"], 2)
        ws = self.load(out).active
        self.assertEqual(ws["A1"].value, "Capabilities Comparison Matrix")
        self.assertEqual(ws["B6"].value, "Weird")
        self.assertEqual(ws["B7"].value, "Medium")
        self.assertEqual(ws["C6"].value, "—")  # platform without 'features' key
        self.assertEqual(ws["B6"].fill.fgColor.rgb, "00FFFFFF")

    def test_formula_injection_neutralized(self):
        out = self.dir / "i.xlsx"
        cfg = {"title": "=HYPERLINK(\"x\")",
               "categories": [{"name": "@cat", "features": [{"name": "+feat", "priority": "-p"}]}],
               "platforms": [{"name": "=P", "features": ["+feat"]}]}
        mb.build_matrix(cfg, str(out))
        ws = self.load(out).active
        self.assertEqual(ws["A1"].value, "'=HYPERLINK(\"x\")")
        self.assertEqual(ws["C2"].value, "'=P")
        self.assertEqual(ws["A5"].value, "'@cat")
        self.assertEqual(ws["A6"].value, "'+feat")
        self.assertEqual(ws["B6"].value, "'-p")
        # the platform's feature list still matches the raw (unsanitized) feature name
        self.assertEqual(ws["C6"].value, "✔")

    def test_clone_style_is_reserved(self):
        out = self.dir / "c.xlsx"
        with self.assertLogs(mb.log, level=logging.WARNING) as logs:
            mb.build_matrix(SAMPLE_CONFIG, str(out), clone_xlsx="whatever.xlsx")
        self.assertIn("not yet implemented", logs.output[0])
        self.assertTrue(out.exists())

    def test_main_and_script(self):
        cfg = self.dir / "cfg.json"
        cfg.write_text(json.dumps(SAMPLE_CONFIG))
        out = self.dir / "o.xlsx"
        code, stdout, _ = call_main(mb, ["--config", str(cfg), "--out", str(out)])
        self.assertEqual(code, 0)
        self.assertEqual(json.loads(stdout)["features"], 4)
        self.assertTrue(out.exists())
        out2 = self.dir / "o2.xlsx"
        code, stdout, _ = run_script(SCRIPTS / "matrix_builder.py",
                                     ["--config", str(cfg), "--out", str(out2), "--clone-style", "x.xlsx"])
        self.assertEqual(code, 0)
        self.assertTrue(out2.exists())


if __name__ == "__main__":
    unittest.main()
